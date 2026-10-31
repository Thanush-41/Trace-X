"""Sprint tracker for TRACE X (SPRINTS.xlsx).

  python tools/tracker.py create                      # build the workbook (refuses to overwrite)
  python tools/tracker.py task T1.3 Done [commit] [note]
  python tools/tracker.py sprint S1 Done [commit] [note]
  python tools/tracker.py req R7 Verified [note]

Updates edit cells in place, so manual edits made in Excel are kept.
Requires openpyxl (pip install openpyxl).
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PATH = Path(__file__).resolve().parent.parent / "SPRINTS.xlsx"
STATUSES = ["Not started", "In progress", "Done", "Blocked", "Cut"]
REQ_STATUSES = ["Not verified", "Verified", "Failing"]

SPRINTS = [
    ("S0", "Repo, plan and tracker", "PLAN.md", 10),
    ("S1", "Scaffold: Next.js, DB schema, triggers, policy, utilities", "§2, §3, §5.1", 20),
    ("S2", "Synthetic data generator + committed seeds 1–3", "§8", 25),
    ("S3", "Ingestion pipeline: parse, aliases, validation, buckets, APIs", "§4, §6", 25),
    ("S4", "Detection engine FD-01/02/03/05/06, scoring, cases + tests", "§5, §9", 50),
    ("S5", "UI part 1: Dashboard, Ingest, Alerts list + detail", "§7.1–7.4", 30),
    ("S6", "UI part 2: Cases + decisions, Customer timeline, Rules, Evaluation", "§7.5–7.8", 30),
    ("S7", "Hardening, README/limitations, deploy, demo rehearsal", "§10, §11", 20),
]

TASKS = [
    ("T0.1", "S0", "git init, first commit, push to GitHub", "main exists on origin"),
    ("T0.2", "S0", "Commit PLAN.md and SPRINTS.xlsx tracker", "Both files on main"),
    ("T1.1", "S1", "create-next-app (TS, App Router, Tailwind), .gitignore", "npm run build passes"),
    ("T1.2", "S1", "Drizzle + libSQL client; schema for source, bookkeeping, detection and case tables", "Tables created on a fresh DB"),
    ("T1.3", "S1", "Immutability triggers on alerts and case_decisions", "UPDATE alerts aborts"),
    ("T1.4", "S1", "config/policy.json with all weights and thresholds", "Loaded and typed by lib/policy.ts"),
    ("T1.5", "S1", "Utilities: normalize, hash, stats (median/MAD/p95), union-find", "Unit tests pass"),
    ("T2.1", "S2", "Seeded PRNG + background population (customers, accounts, txns, logins, loans)", "Deterministic output per seed"),
    ("T2.2", "S2", "Planted fraud scenarios S1–S6", "Present in ground_truth.json"),
    ("T2.3", "S2", "Planted lookalikes L1–L4", "Present in ground_truth.json"),
    ("T2.4", "S2", "Outputs: CSV ×5, XLSX, JSON, errors file, alias-header file, ground truth", "Seeds 1–3 committed"),
    ("T3.1", "S3", "Parsers: CSV, XLSX (multi-sheet), JSON (array or keyed)", "Same seed → same rows in all 3 formats"),
    ("T3.2", "S3", "Header aliases + dataset type detection + override", "Alias file maps correctly"),
    ("T3.3", "S3", "zod row validation + value normalisation", "Bad rows report field + message"),
    ("T3.4", "S3", "Buckets: accepted / rejected / duplicate / warning; conflict detection", "total = sum of buckets"),
    ("T3.5", "S3", "APIs: /api/ingest, /api/samples, /api/samples/load, /api/batches, errors.csv", "Ingestion tests pass"),
    ("T4.1", "S4", "DetectionContext: indexes, sorted histories, baselines", "Built once per run"),
    ("T4.2", "S4", "FD-01 suspicious transfer", "S2–S4 detected; L2 not Medium+"),
    ("T4.3", "S4", "FD-02 shared device", "S5 device High; L1 Low"),
    ("T4.4", "S4", "FD-03 account takeover", "S1 High; L3 Low; L4 none"),
    ("T4.5", "S4", "FD-05 loan application ring", "S5 High"),
    ("T4.6", "S4", "Scoring, floors, fingerprints, insert-only persistence", "Re-run adds 0 alerts"),
    ("T4.7", "S4", "FD-06 cross-system correlation + hub exclusion", "S5+S6 in one component"),
    ("T4.8", "S4", "Case builder + /api/detect", "Re-run adds 0 cases"),
    ("T4.9", "S4", "Detection + integrity tests on seeds 1–3", "vitest green"),
    ("T5.1", "S5", "App shell, nav, shared UI components", "All routes reachable"),
    ("T5.2", "S5", "Dashboard + /api/dashboard", "Counts match DB"),
    ("T5.3", "S5", "Ingest screen: sample picker, upload, result cards, history", "Bad file shows 5 issues"),
    ("T5.4", "S5", "Alerts list: tabs + 6 filters + /api/alerts", "Each filter narrows results"),
    ("T5.5", "S5", "Alert detail: primary event, why-flagged table, source rows", "Every evidence ref resolves"),
    ("T6.1", "S6", "Case queue + case detail + merged timeline", "FD-06 case shows all members"),
    ("T6.2", "S6", "Decision form + history + revised level/category", "Persists after refresh; alert unchanged"),
    ("T6.3", "S6", "Customer investigation page + timeline", "Linked devices list other customers"),
    ("T6.4", "S6", "Rules page", "Shows policy version + weights"),
    ("T6.5", "S6", "Evaluation page + /api/evaluate", "All seed-3 scenarios pass"),
    ("T7.1", "S7", "Reset workspace + error handling pass", "No stack traces reach the client"),
    ("T7.2", "S7", "README + KNOWN_LIMITATIONS", "Run steps work from a clean clone"),
    ("T7.3", "S7", "Deploy (Vercel + Turso)", "Deployed URL passes §11"),
    ("T7.4", "S7", "Demo rehearsal on seed 3", "§11 completed in ≤ 4 min"),
]

REQS = [
    ("R1", "Upload or select data from different systems", "Ingest screen: upload + sample picker"),
    ("R2", "Customer/account, transaction, device/login data", "customers, accounts, transactions, login_events"),
    ("R3", "CSV, Excel, JSON", "lib/ingest/parse.ts"),
    ("R4", "Validate; flag missing mandatory fields; never silently ignore", "Validation buckets + error report"),
    ("R5", "Common structure linked by shared IDs", "Schema + application_id on disbursements"),
    ("R6", "Configurable rules + statistics", "config/policy.json + Rules page"),
    ("R7", "Suspicious transfers", "FD-01"),
    ("R8", "Shared devices and whether suspicious", "FD-02"),
    ("R9", "Repeated login failures → success / sensitive activity", "FD-03"),
    ("R10", "Legitimate lookalikes", "Mitigations + Reviewed tab + L1–L4"),
    ("R11", "At least one extra scenario", "FD-05 (+ FD-06)"),
    ("R12", "Multi-indicator risk classification", "lib/detect/score.ts"),
    ("R13", "Normal high value ≠ fraud; after login risk scores higher", "FD-01 WITHIN_CUSTOMER_NORM / PRECEDED_BY_LOGIN_RISK"),
    ("R14", "Explain contributing factors", "Why-flagged table"),
    ("R15", "Dashboard metrics", "Dashboard screen"),
    ("R16", "Filters: customer, transaction, product, risk, date, detector", "Alerts screen + /api/alerts"),
    ("R17", "Alert shows event, reasons, supporting records", "Alert detail"),
    ("R18", "Case: ID, customer/account, category, risk, evidence, status", "cases + case_decisions"),
    ("R19", "Record a decision", "Case decision form"),
    ("R20", "Preserve original alert after reclassification", "Insert-only alerts + trigger"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_COLORS = {
    "Done": "C6EFCE", "Verified": "C6EFCE",
    "In progress": "FFEB9C",
    "Blocked": "FFC7CE", "Failing": "FFC7CE",
    "Cut": "D9D9D9",
}


def _sheet(wb, title, headers, widths, rows, status_col, statuses):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    col = ws.cell(1, status_col).column_letter
    last = ws.max_row
    dv = DataValidation(type="list", formula1='"' + ",".join(statuses) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{last}")
    span = f"A2:{ws.cell(1, len(headers)).column_letter}{last}"
    for status, color in STATUS_COLORS.items():
        if status in statuses:
            ws.conditional_formatting.add(
                span,
                FormulaRule(formula=[f'${col}2="{status}"'], fill=PatternFill("solid", fgColor=color)),
            )
    return ws


def create():
    if PATH.exists():
        sys.exit(f"{PATH.name} already exists; use the update commands instead.")
    wb = Workbook()
    wb.remove(wb.active)

    ws = _sheet(
        wb, "Sprints",
        ["Sprint", "Goal", "Plan sections", "Timebox (min)", "Tasks", "Tasks done", "Status", "Completed on", "Commit", "Notes"],
        [8, 52, 16, 13, 8, 11, 13, 14, 12, 40],
        [(s, g, p, t, None, None, "Not started", None, None, None) for s, g, p, t in SPRINTS],
        7, STATUSES,
    )
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 5).value = f'=COUNTIF(Tasks!$B:$B,A{r})'
        ws.cell(r, 6).value = f'=COUNTIFS(Tasks!$B:$B,A{r},Tasks!$E:$E,"Done")'
    total = ws.max_row + 1
    ws.cell(total, 1).value = "Total"
    ws.cell(total, 1).font = Font(bold=True)
    ws.cell(total, 4).value = f"=SUM(D2:D{total - 1})"
    ws.cell(total, 5).value = f"=SUM(E2:E{total - 1})"
    ws.cell(total, 6).value = f"=SUM(F2:F{total - 1})"

    _sheet(
        wb, "Tasks",
        ["ID", "Sprint", "Task", "Done when", "Status", "Owner", "Completed on", "Commit", "Notes"],
        [7, 8, 60, 38, 13, 12, 14, 12, 40],
        [(i, s, t, d, "Not started", None, None, None, None) for i, s, t, d in TASKS],
        5, STATUSES,
    )
    _sheet(
        wb, "Requirements",
        ["ID", "Official requirement", "Implemented by", "Status", "Verified on", "Notes"],
        [6, 55, 45, 14, 14, 40],
        [(i, r, b, "Not verified", None, None) for i, r, b in REQS],
        4, REQ_STATUSES,
    )
    wb.save(PATH)
    print(f"created {PATH}")


def _update(sheet, key, status, extra, commit_col, date_col, notes_col, statuses):
    if status not in statuses:
        sys.exit(f"status must be one of {statuses}")
    wb = load_workbook(PATH)
    ws = wb[sheet]
    status_col = [c.value for c in ws[1]].index("Status") + 1
    for row in ws.iter_rows(min_row=2):
        if row[0].value == key:
            r = row[0].row
            ws.cell(r, status_col).value = status
            if status in ("Done", "Verified"):
                ws.cell(r, date_col).value = date.today().isoformat()
            if commit_col and extra:
                ws.cell(r, commit_col).value = extra.pop(0)
            if extra:
                ws.cell(r, notes_col).value = extra.pop(0)
            wb.save(PATH)
            print(f"{sheet} {key} → {status}")
            return
    sys.exit(f"{key} not found in {sheet}")


if __name__ == "__main__":
    cmd, *args = sys.argv[1:] or ["help"]
    if cmd == "create":
        create()
    elif cmd == "task" and len(args) >= 2:
        _update("Tasks", args[0], args[1], args[2:], 8, 7, 9, STATUSES)
    elif cmd == "sprint" and len(args) >= 2:
        _update("Sprints", args[0], args[1], args[2:], 9, 8, 10, STATUSES)
    elif cmd == "req" and len(args) >= 2:
        _update("Requirements", args[0], args[1], args[2:], None, 5, 6, REQ_STATUSES)
    else:
        print(__doc__)
